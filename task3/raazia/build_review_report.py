from pathlib import Path
import json,csv,datetime
from docx import Document
from docx.shared import Inches,Pt,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parents[1];out=root/'raazia';d=json.loads((out/'review_data.json').read_text());overlay=json.loads((root/'logs/reviewed_data.json').read_text())
# Explicit eligibility appraisal. Publication identity and final-version equivalence are distinct.
rows=[]
mechanisms={'P008':'MPS allocation among concurrent inference functions','P009':'Temporal scheduling of multiple inference models on one GPU','P010':'Single-GPU inference/training co-location with separable inference P99','P011':'MPS processes within a MIG instance','P012':'Many low-rate inference functions temporally share one GPU','P013':'Streams and runtime/API sharing across inference workflows','P014':'Kernel interception and temporal/MIG sharing with separate inference measurements'}
for s in d['sources']:
 p=s['Report ID'];qa=[q for q in d['qa'] if q['Report ID']==p];es=[e for e in d['outcomes'] if e['Report ID']==p]
 checks=[('IC1','Supported by accessed conference copy' if p!='P011' else 'Final-version condition open','Conference title/authors/year matched in full text.' if p!='P011' else 'SC24/author records confirm peer-reviewed publication, but numerical source is arXivv1 and final-text equivalence is not established.',('PDF title/cover page' if p!='P011' else 'SC24 official record; author publication page; arXivv1 title page')),('IC2','Supported',mechanisms[p],qa[1]['Source anchor']),('IC3','Supported','Inference workload is explicit; training/rendering outcomes are kept separate where present.',qa[7]['Source anchor']),('IC4','Supported','Retained outcome values and methods are inspectable in the accessed source.','; '.join(e['Evidence location'] for e in es)),('IC5','Supported','Extracted evidence concerns co-tenancy or control on one physical GPU; fleet-only outcomes not substituted.',qa[7]['Source anchor']),('IC6','Supported',str(s['Year'])+' is within the submitted2020-2026 window.','Conference title page / source identity')]
 for criterion,status,basis,anchor in checks:rows.append({'Study ID':s['Study ID'],'Report ID':p,'Criterion':criterion,'Review finding':status,'Basis':basis,'Anchor':anchor,'Reviewer':d['reviewer'],'Checked UTC':d['checked_utc']})
with (out/'Eligibility_Source_Review.csv').open('w',newline='') as f:w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
d['eligibility']=rows;(out/'review_data.json').write_text(json.dumps(d,indent=2))
doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Cm(2);sec.left_margin=sec.right_margin=Cm(2.2)
for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
 st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.space_after=Pt(6)
for name in ['Heading 1','Heading 2','Heading 3']:doc.styles[name].paragraph_format.keep_with_next=True
def para(t):return doc.add_paragraph(t)
def table(headers,rs,widths=None):
 t=doc.add_table(rows=1, cols=len(headers));t.autofit=False
 for i,h in enumerate(headers):t.rows[0].cells[i].text=h
 for rr in rs:
  for c,v in zip(t.add_row().cells,rr):c.text=str(v)
 for ri,row in enumerate(t.rows):
  if ri==0:
   hdr=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(hdr)
  for ci,c in enumerate(row.cells):
   if widths:c.width=Cm(widths[ci])
   pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    n=OxmlElement('w:'+edge);n.set(qn('w:val'),'single');n.set(qn('w:sz'),'4');n.set(qn('w:color'),'D9D9D9');b.append(n)
   pr.append(b)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EDF2');pr.append(sh)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(4)
    for run in p.runs:run.font.size=Pt(10);run.bold=(ri==0)
 if widths:
  for col,w in zip(t.columns,widths):col.width=Cm(w)
 return t

doc.add_heading('Raazia Assigned Source Review',0)
para('Operating Systems CEP Topic 23 | Studies S008 to S014 | 21 October 2026')
para('This deliverable records an AI-assisted source review completed under Raazia Imran Reshamwala’s direction. It contains seven source identity checks, 42 eligibility-criterion assessments, 70 QA appraisals, 13 outcome checks, a correction register and reviewed bibliography. The integrated report and master workbook include the supported corrections.')
para('All 13 retained numerical values were found in the accessed sources. Corrections concern page anchors, software details, runtime-mode distinctions, offered-load interpretation and claim-specific QA rationale. Original scores and records remain in the frozen baseline.')
table(['Study','Original QA','Reviewed QA','Outcome rows'],[[s['Study ID']+' '+s['Title'].split(':')[0],s['Original QA'],s['Reviewed QA'],sum(e['Report ID']==s['Report ID'] for e in d['outcomes'])] for s in d['sources']],[8.7,2.3,2.3,2.3])
para('QA totals change from 110/140 to 114/140. Clockwork and StreamBox each gain two points because reported configuration and clearly delimited performance/trust evidence satisfy Q4 and Q9. A score is a source appraisal, not a performance measure or a human independent vote.')
para('Version boundary: ParvaGPU values were checked against arXiv2409.14447v1. Conference publication identity is confirmed; equivalence of the numerical manuscript to the final published full text is not established. Global review scope is unchanged: 559 retained records have not received source assessment. These are study-selection limitations, not findings that this seven-study check can erase.')
doc.add_heading('Integrated corrections',1)
for c in d['changes']:
 doc.add_heading(c['Change ID']+' '+c['Scope']+' '+c['Field'],2);para('Original: '+c['Original']);para('Reviewed treatment: '+c['Reviewed treatment']);para('Evidence: '+c['Evidence'])
for s in d['sources']:
 doc.add_page_break();pid=s['Report ID'];doc.add_heading(s['Study ID']+' '+s['Title'],1)
 para('Report '+pid+' | '+str(s['Year'])+' | '+s['Venue']);para('DOI or proceedings: '+s['DOI or proceedings']);para('Source: '+s['Source URL']);para('Version checked: '+s['Version checked']);para('PDF pages: '+str(s['PDF pages'])+' | SHA256: '+s['SHA256'])
 doc.add_heading('Eligibility appraisal',2)
 table(['Criterion','Finding and basis'],[[x['Criterion'],x['Review finding']+'. '+x['Basis']+' Anchor: '+x['Anchor']] for x in rows if x['Report ID']==pid],[2,13.6])
 doc.add_heading('Quality appraisal',2)
 table(['Item','Original','Reviewed','Rationale and source anchor'],[[q['Criterion'],q['Original score'],q['Reviewed score'],q['Rationale']+' '+q['Source anchor']] for q in d['qa'] if q['Report ID']==pid],[1.2,1.7,1.7,11])
 doc.add_heading('Verified reported outcomes',2)
 for e in d['outcomes']:
  if e['Report ID']!=pid:continue
  para(e['Evidence ID']+' '+e['Metric']+': '+str(e['Value'])+' '+e['Units']+'.')
  para('Baseline: '+e['Baseline']+'. Workload: '+e['Workload']+'. Hardware: '+e['Hardware']+'.')
  para('Interpretation: '+e['Interpretation']+' Anchor: '+e['Evidence location']+'.')
para('Review attribution: '+d['reviewer']+'. Checked UTC: '+d['checked_utc']+'. Source-check records do not assert a second independent human review or experimental reproduction.')
for sec in doc.sections:
 p=sec.footer.paragraphs[0];p.alignment=1;f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)

for tree in [doc._element,doc.styles.element]:
 for n in list(tree.iter(qn('w:pBdr'))):n.getparent().remove(n)
doc.save(out/'Raazia_Assigned_Review.docx');print(out/'Raazia_Assigned_Review.docx')
