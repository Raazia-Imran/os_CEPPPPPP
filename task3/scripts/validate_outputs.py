from pathlib import Path
import openpyxl,json,collections,hashlib,datetime,subprocess
from zoneinfo import ZoneInfo
root=Path(__file__).resolve().parents[1]
a=openpyxl.load_workbook(root/'baseline/CEP_EVIDENCE_WORKBOOK.xlsx',read_only=True,data_only=True)
b=openpyxl.load_workbook(root/'workbook/CEP_Task3_Evidence_Workbook.xlsx',read_only=True,data_only=True)
diffs=[]
for s in a:
 aa=list(s.values);bb=list(b[s.title].values)
 for i,row in enumerate(aa):
  for j,x in enumerate(row):
   y=bb[i][j] if i<len(bb) and j<len(bb[i]) else None
   if x!=y:
    if s.title=='QA' and j in [2,3] and x is not None and str(x)==str(y):continue
    diffs.append([s.title,i+1,j+1,str(x),str(y)])
checks={}
checks['original_sheet_value_differences_excluding_numeric_QA_conversion']=diffs
checks['sheet_count']=len(b.sheetnames)
checks['count_reconciliation_differences']=[r[3] for r in list(b['Count Reconciliation'].values)[1:]]
checks['QA_reconciliation_differences']=[list(r[8:10]) for r in list(b['QA Reconciliation'].values)[1:]]
checks['source_scores_below_60_percent']=sum(float(r[8])<.6 for r in list(a['Studies'].values)[1:])
checks['pdf_info']=subprocess.check_output(['pdfinfo',str(root/'report/main.pdf')],text=True)
checks['pdf_sha256']=hashlib.sha256((root/'report/main.pdf').read_bytes()).hexdigest()
checks['workbook_sha256']=hashlib.sha256((root/'workbook/CEP_Task3_Evidence_Workbook.xlsx').read_bytes()).hexdigest()
checks['latex_material_warnings']=[l for l in (root/'report/main.log').read_text().splitlines() if any(t in l for t in ['Overfull','undefined','Missing character'])]
checks['checked_utc']=datetime.datetime.now(datetime.timezone.utc).isoformat()
checks['scope']='Technical preservation and arithmetic checks; not independent scientific verification'
(root/'logs/validation.json').write_text(json.dumps(checks,indent=2));print(json.dumps({k:v for k,v in checks.items() if k not in ['pdf_info','QA_reconciliation_differences']},indent=2))
