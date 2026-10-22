from pathlib import Path
import json,sys
from docx import Document
from docx.shared import Cm,Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
out=Path(sys.argv[1]);d=json.loads((out/'review_data.json').read_text());doc=Document();s=doc.sections[0];s.top_margin=s.bottom_margin=Cm(1.7);s.left_margin=s.right_margin=Cm(1.8)
for name in ['Normal','Title','Heading 1','Heading 2']:doc.styles[name].font.name='Calibri'
doc.styles['Normal'].font.size=Pt(10);doc.styles['Normal'].paragraph_format.space_after=Pt(5)
def p(t):doc.add_paragraph(str(t))
def table(h,rr,widths):
 t=doc.add_table(rows=1,cols=len(h));t.style='Light Shading Accent 1';t.autofit=False
 for col,w in zip(t.columns,widths):col.width=Cm(w)
 for c,v in zip(t.rows[0].cells,h):c.text=str(v)
 hdr=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(hdr)
 for row in rr:
  for c,v in zip(t.add_row().cells,row):c.text=str(v)
 for row in t.rows:
  z=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(z)
  for c,w in zip(row.cells,widths):
   c.width=Cm(w)
   for pp in c.paragraphs:
    pp.paragraph_format.space_after=Pt(3)
    for r in pp.runs:r.font.size=Pt(9)
doc.add_heading('Eshaal Assigned Source Review',0);p('Operating Systems CEP Topic 23 | 22 October 2026')
p('Eight report copies, seven unique studies. First source review prepared at Raazia’s request; human independent signoff remains pending.')
table(['Item','Reviewed','Remaining'],[['Primary-study QA items',70,'Independent adjudication'],['Report eligibility criteria',48,'Human independent checks'],['Retained outcomes',30,'Independent source checks'],['Methodology audit findings',14,'Coverage/provenance issues listed below']],[8,3,6.4])
p('Original QA scores and frozen baseline are preserved. The source-based review verifies retained values and identifies interpretation/version corrections. It is not a rerun of experiments or a complete census of all screening decisions.')
p('The three packages now cover all20included studies. Master integration is complete. The20study target is met; independent human checks,559record selection and historical search limitations remain open.')
doc.add_heading('Main corrections',1)
for c in d['changes']:p(c['Change ID']+' / '+c['Scope']+': '+c['Proposed treatment']+'. Source: '+c['Evidence'])
for src in d['sources']:
 pid=src['Report ID'];heading=doc.add_heading(src['Study ID']+' / '+pid+' - '+src['Title'],1);heading.paragraph_format.page_break_before=True
 p('Version: '+src['Version checked']);p('Source: '+src['Source URL']);p('DOI/proceedings: '+src['DOI']);p('PDF pages: '+str(src['PDF pages'])+'. SHA256: '+src['SHA256'])
 er=[r for r in d['eligibility'] if r['Report ID']==pid]
 table(['Criterion','Finding and source anchor'],[[r['Item ID'],r['Rationale']+' '+r['Source anchor']] for r in er],[1.8,15.6])
 qr=[r for r in d['qa'] if r['Report ID']==pid]
 if qr:
  doc.add_heading('Quality appraisal',2);p('Original QA: '+str(src['Original QA'])+'/'+str(src['Original maximum'])+'; current appraised QA: '+str(src['Reviewed QA'])+'/'+str(src['Reviewed maximum'])+'. Source-based calibration is logged; independent adjudication remains pending.')
  table(['Item','Score','Rationale and anchor'],[[r['Item ID'],'N/A' if r['Original score']=='NA' else str(r['Reviewed score']),r['Rationale']+' '+r['Source anchor']] for r in qr],[1.2,1.4,14.8])
 else:p('Companion report linked toS005. No extra study-level QA form or new study count is created. Its new NVIDIA/LLM experiments must be extracted as separate report/version strata if used quantitatively.')
 es=[r for r in d['outcomes'] if r['Report ID']==pid]
 if es:
  doc.add_heading('Checked retained outcomes',2)
  table(['Evidence','Value and metric','Context and anchor'],[[r['Evidence ID'],str(r['Value'])+' '+r['Units']+'; '+r['Metric'],r['Hardware']+'. '+r['Workload']+'. Baseline: '+r['Baseline']+'. '+r['Evidence location']] for r in es],[1.5,5,10.9])
  for text in dict.fromkeys(r['Review note'] for r in es):p(text)
 if pid=='P023':p('The first-page extension statement links P023 to P005. Source setup: TVM0.14/Ubuntu20.04/32GB host and full NVIDIA implementation, versus conference TVM0.8/Ubuntu18.04/16GB host and restrictedREEF-N. Do not transplant journal performance into the conferenceAMD outcome rows. Journal Figures28-30 add Triton,Orion and LLM comparisons under distinct conditions.')
doc.add_page_break();doc.add_heading('Methodology audit and actions',1)
for r in d['methodology']:
 doc.add_heading(r['Audit ID']+' '+r['Topic']+' ('+r['Status']+')',2);p(r['Finding']);p('Evidence: '+r['Evidence']);p('Next: '+r['Next action'])
doc.add_heading('Independent-check return',1);p('Rameen independently checks P001–P007 plusP023. Eshaal independently checks Raazia’sP008–P014. Enter the actual reviewer,checked time,initial decision,source anchor and disagreements separately. These prepared review records cannot be copied into human-vote columns as independent signoff.')
p('Checked: '+d['built_pkt']+'. Attribution: '+d['reviewer']+'. Full texts are excluded from public Git; source hashes and appraisal facts are retained.')
for tree in [doc._element,doc.styles.element]:
 for n in list(tree.iter(qn('w:pBdr'))):n.getparent().remove(n)
doc.save(out/'Eshaal_Assigned_Review.docx');print('Saved Eshaal_Assigned_Review.docx')
