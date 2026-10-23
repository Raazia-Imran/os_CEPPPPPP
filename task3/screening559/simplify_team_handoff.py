"""Produce a one-page handoff and count-based progress artifacts."""
from pathlib import Path
import csv,json,shutil,hashlib
from docx import Document
from docx.shared import Inches,Pt,RGBColor
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Image,Table,TableStyle,PageBreak
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
r=Path(__file__).resolve().parents[2];d=r/'task3/screening559';s=r/'Task3_First_Draft_2026-10-24';t=s/'04_Team';progress=json.loads((d/'progress.json').read_text());closed=progress['final_source_based_eligibility_decisions'];op=334-closed
text=f'''# What your teammates need to do

Use CEP_Task3_Team_Package_2026-10-24.zip and CEP_Task3_LaTeX_Project_2026-10-24.zip. The older two ZIPs are backups. Both current packages describe the same 30 studies and 31 reports.

## Eshaal edits LaTeX
1. Upload the LaTeX ZIP to Overleaf. Set main.tex as the main document and use pdfLaTeX.
2. Read the report from Introduction to Conclusion. Improve clear wording, headings, table wrapping and spelling. Keep references and evidence appendices.
3. Keep claims, numbers, paper IDs and the stated research limitations consistent with the master workbook. Ask Rameen to verify any factual change.
4. Recompile. Return the edited project ZIP and matching PDF as one version.

## Rameen checks evidence
1. Open 01_Report/CEP_Task3_First_Draft.pdf and 02_Master/CEP_Task3_Evidence_Workbook.xlsx from the team package.
2. Use 04_Team/New_Paper_Map.csv to locate the ten added papers and their PDFs in 05_Added_Paper_PDFs. Check that each cited result has the same metric, value, workload and source page as the workbook. Read the QA and eligibility evidence for those papers.
3. Check the report counts: 30 studies, 31 reports, 300 QA items, 90 reported observations, two separate calculations; continuation eligibility {closed}/334 closed and {op} open. Confirm citations appear in the reference list. Reuse verified DOIs; investigate a DOI only if a mismatch appears.
4. Send Eshaal a short corrections list: paper ID, report location, source page, old wording/number and corrected wording/number. Record your name/date only for checks actually completed.

## Finish together
Rameen checks Eshaal's rebuilt PDF against the corrections list. Both confirm the cover details, references, tables and all 13 report sections. Eshaal returns one final PDF and one matching LaTeX ZIP; retain the matching Word/master/logs in the team package. No additional screening task is assigned to Raazia by this guide.

## Research still open
The current draft is assembled; the entire SLR is incomplete. {op} advanced candidates still need identity/version and eligibility decisions. Of the original 320 open cases, 82 had indexed PDF routes attempted this cycle; 34 PDFs were recovered and two records were excluded after source assessment. The remaining {op} comprise 32 recovered texts awaiting decisions, 48 cases without a recovered PDF from those routes, and 238 cases needing further source discovery. These categories do not justify automatic exclusions. The separate access log records the attempts. Genuine independent checking and historical search/PRISMA reconciliation also remain.
'''
(t/'READ_THIS_FIRST.md').write_text(text)
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.6);sec.left_margin=sec.right_margin=Inches(.65)
for name in ['Normal','Title','Heading 1','Heading 2']:
 st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(10);doc.styles['Normal'].paragraph_format.space_after=Pt(4)
doc.styles['Title'].font.size=Pt(19);doc.styles['Title'].paragraph_format.space_after=Pt(8)
doc.styles['Heading 1'].font.size=Pt(12);doc.styles['Heading 1'].paragraph_format.space_before=Pt(8);doc.styles['Heading 1'].paragraph_format.space_after=Pt(4)
for line in text.splitlines():
 if not line:continue
 if line.startswith('# '):doc.add_paragraph(line[2:],'Title')
 elif line.startswith('## '):doc.add_paragraph(line[3:],'Heading 1')
 else:doc.add_paragraph(line)
from docx.oxml.ns import qn
for st in doc.styles:
 pp=st.element.find(qn('w:pPr'))
 if pp is not None:
  for border in list(pp.findall(qn('w:pBdr'))):pp.remove(border)
doc.save(t/'Team_Guide.docx')
# Archive the longer optional check queues to keep the starting view simple.
a=t/'Detailed_Audit_Queues';a.mkdir(exist_ok=True)
for name in ['Independent_Check_Queue.csv','Eshaal_Checks.csv','Rameen_Checks.csv']:
 if (t/name).exists():shutil.move(t/name,a/name)
(t/'Team_Correction_Inbox.csv').write_text('Checker,Date,Paper ID,Report location,Source page,Current wording or value,Correction,Reason\n')
# Source access and new decisions accompany the submitted logs.
for name in ['Public_Full_Text_Access_Attempts_2026-10-24.csv','eligibility_cycle5/Final_Exclusions.csv','eligibility_cycle5/Eligibility.csv','eligibility_cycle5/Queue_Access_Status.csv']:
 dest=s/'03_Data_and_Logs'/('Cycle5_'+Path(name).name if '/' in name else name);shutil.copyfile(d/name,dest)
# Measurable phases have explicit denominators; no invented overall percentage.
phases=[('Canonical records with a closed status',4075-op,4075,'Includes screening exclusions; not overall SLR completion'),('Continuation screening decisions',559,559,'Title/abstract stage; title-only evidence explicitly retained'),('Continuation eligibility decisions',closed,334,'10 included and 6 full-text exclusions'),('Added included PDF copies verified',10,10,'Only the ten additions; not all advanced records'),('Included-study QA items recorded',300,300,'Current 30-study corpus; independent checks separate'),('Included outcome rows integrated',90,90,'Current 30-study corpus; source observations, not pooled effects'),('Main report sections present',13,13,'Presence only; not final research compliance'),('Required deliverables mapped',25,25,'Mapped to files; not all conditions closed')]
with (t/'Phase_Progress.csv').open('w',newline='') as f:
 w=csv.writer(f);w.writerow(['Phase','Done','Denominator','Percent','Meaning']);w.writerows((n,v,total,round(v/total*100,2),note) for n,v,total,note in phases)
fig,ax=plt.subplots(figsize=(11,5.4));fig.patch.set_facecolor('white');labels=[p[0] for p in phases];vals=[p[1]/p[2]*100 for p in phases];ys=list(range(len(phases)))
ax.barh(ys,[100]*len(ys),color='#e6ebf0',height=.58);ax.barh(ys,vals,color=['#357b9b' if v==100 else '#c27c23' for v in vals],height=.58)
ax.set_yticks(ys,labels,fontsize=10);ax.invert_yaxis();ax.set_xlim(0,116);ax.set_xticks([0,25,50,75,100],['0%','25%','50%','75%','100%']);ax.tick_params(axis='y',length=0)
for i,(v,p) in enumerate(zip(vals,phases)):ax.text(101,i,f'{v:.1f}%  ({p[1]}/{p[2]})',va='center',fontsize=9)
for sp in ax.spines.values():sp.set_visible(False)
ax.set_title('Recorded progress by phase',loc='left',fontsize=17,pad=18);fig.text(.02,.015,'Each bar has its own denominator. These percentages must not be averaged into an overall completion score.',fontsize=10);fig.tight_layout(rect=(0,.045,1,1));fig.savefig(t/'Progress.png',dpi=180);plt.close(fig)
headings=[(1,'Background','Drafted'),(2,'Engineering problem','Drafted'),(3,'Objective','Drafted'),(4,'SLR methodology','Documented with open selection'),(5,'Initial activities and RQs','Drafted'),(6,'Protocol','Preserved; later amendments separate'),(7,'Databases','Documented; coverage deviations remain'),(8,'Search strategy','Recorded; historical run mapping incomplete'),(9,'Literature period','Specified'),(10,'Eligibility criteria','Specified; application incomplete'),(11,'PRISMA selection','Open; 318 candidates unresolved'),(12,'Search and screening log','Recorded; eligibility still open'),(13,'Quality assessment','300 items recorded for included set'),(14,'Data extraction','30 study forms and 90 outcomes'),(15,'Taxonomy','Integrated for assessed set'),(16,'Comparison','Integrated with comparability limits'),(17,'Engineering trade-offs','Integrated for assessed set'),(18,'Conflicting evidence','Included; adverse results retained'),(19,'RQ synthesis','All five RQs drafted'),(20,'Recommendation','Conditional recommendations drafted'),(21,'Research gaps','Derived from assessed evidence'),(22,'Validity','Limits documented'),(23,'Mandatory deliverables','25 mapped; final closure incomplete'),(24,'Report structure','13/13 main sections present'),(25,'Analytical standard','Integrated for assessed set'),(26,'Research quality','Traceability present; selection/provenance open'),(27,'Engineering answer','Drafted; bounded by assessed set')]
with (t/'All_Guideline_Headings.csv').open('w',newline='') as f:
 w=csv.writer(f);w.writerow(['Heading','Requirement','Status']);w.writerows(headings)
styles=getSampleStyleSheet();story=[Paragraph('Task 3 progress and remaining research',styles['Title']),Paragraph(f'30 studies across 31 reports. Continuation eligibility: {closed}/334 closed, {op} open. Whole-review selection and historical search reconciliation remain incomplete.',styles['BodyText']),Spacer(1,12),Image(str(t/'Progress.png'),width=520,height=255),Spacer(1,12)]
for n,v,total,note in phases:story.append(Paragraph(f'<b>{n}</b>: {v}/{total} ({v/total*100:.1f}%). {note}.',styles['BodyText']));story.append(Spacer(1,5))
story += [PageBreak(),Paragraph('Status against every guideline heading',styles['Title']),Paragraph('Statuses describe the recorded draft. Percentages are omitted when no objective denominator exists. Drafted does not mean independently verified or fully compliant.',styles['BodyText']),Spacer(1,10)]
data=[['No','Requirement','Current status']]+[[str(a),Paragraph(b,styles['BodyText']),Paragraph(c,styles['BodyText'])] for a,b,c in headings]
table=Table(data,colWidths=[28,175,315],repeatRows=1);table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#173b50')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),.4,colors.HexColor('#d9d9d9')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4)]));story.append(table)
SimpleDocTemplate(str(t/'Progress.pdf'),pagesize=(612,792),leftMargin=46,rightMargin=46,topMargin=36,bottomMargin=36).build(story)
(s/'START_HERE.md').write_text('# Current first draft\n\nStart with 04_Team/READ_THIS_FIRST.md or the one-page Team_Guide.pdf. Eshaal owns the LaTeX edit; Rameen checks evidence and returns corrections. The older frozen ZIPs are backups.\n\nCurrent corpus: 30 studies / 31 reports, 300 QA items, 90 observations and two separate calculations. Continuation eligibility: 16/334 closed; 318 open. The entire SLR is incomplete. See 04_Team/Progress.pdf for actual denominators and every guideline heading.\n\n01_Report holds PDF/Word; 02_Master holds Excel; 03_Data_and_Logs holds CSVs and decisions; 05_Added_Paper_PDFs contains ten private reading copies in the ZIP; 06_Protocol_and_Guidelines holds requirements; 07_Historical_Backup preserves the original release. 08_LaTeX in the full ZIP contains the self-contained editable project. Private retrieval reading copies are in 09_Eligibility_Reading, with open status explicit. These private PDFs are absent from public GitHub.\n')
print('One-page guide and progress artifacts generated')
