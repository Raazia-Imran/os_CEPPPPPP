"""Create the editable Word companion from the compiled LaTeX source and citation order."""
from pathlib import Path
import re, subprocess, json, os
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parents[2]; report=Path(__file__).resolve().parent
tmp=Path('/tmp/task3_integrated_word');tmp.mkdir(exist_ok=True)
def expand(text):
    return re.sub(r'\\input\{([^}]+)\}',lambda m:expand((report/m[1]).read_text()),text)
s=expand((report/'main.tex').read_text())
s=s.split(r'\begin{document}',1)[1].split(r'\end{document}',1)[0]
bbl=(report/'main.bbl').read_text()
refs=re.findall(r'\\bibitem\{([^}]+)\}\s*(.*?)(?=\\bibitem|\\end\{thebibliography\})',bbl,re.S)
numbers={key:str(i+1) for i,(key,_) in enumerate(refs)}
s=re.sub(r'\\cite\{([^}]+)\}',lambda m:'['+', '.join(numbers[k] for k in m[1].split(','))+']',s)
s=re.sub(r'\\bibliographystyle\{[^}]+\}', '',s)
reftext='\\section*{References}\n'+ '\n\n'.join('['+str(i+1)+'] '+v.replace(r'\newblock',' ') for i,(_,v) in enumerate(refs))
s=s.replace(r'\bibliography{references}',reftext)
s=re.sub(r'\\nolinkurl\{([^}]+)\}',lambda m:m[1].replace('_',r'\_'),s)
s=s.replace(r'\begin{abstract}',r'\section*{Abstract}').replace(r'\end{abstract}','')
s=s.replace(r'\tableofcontents',r'\section*{Contents}'+'\nSee the Word navigation pane for the complete section hierarchy.\n')
s=s.replace('figures/Frozen_PRISMA.pdf',str(root/'baseline/PRISMA_2020_FLOW_DIAGRAM.png'))
aux=(report/'main.aux').read_text()
labels=dict(re.findall(r'\\newlabel\{([^}]+)\}\{\{([^}]+)\}',aux))
s=re.sub(r'\\ref\{([^}]+)\}',lambda m:labels.get(m[1],m[1]),s)
s=s.replace(r'\addlinespace','')
s=re.sub(r'(&\s*)([0-9]+(?:\.[0-9]+)?)',r'\1{\2}',s)
s=re.sub(r'\\url\{([^}]+)\}',lambda m:r'\url{'+m[1].replace('%',r'\%')+'}',s)
(tmp/'expanded.tex').write_text(s)
subprocess.run(['pandoc',str(tmp/'expanded.tex'),'-f','latex','-t','docx','-o',str(tmp/'converted.docx')],check=True,cwd=report)
doc=Document(tmp/'converted.docx')
expected_tables=s.count(r'\begin{longtable}')
assert len(doc.tables)==expected_tables, f'Table conversion incomplete: {len(doc.tables)} vs {expected_tables}'
for section in doc.sections:
    section.page_width=Mm(210);section.page_height=Mm(297)
    section.top_margin=section.bottom_margin=Mm(22)
    section.left_margin=section.right_margin=Mm(24)
for style in doc.styles:
    if style.type==1:
        style.font.name='Times New Roman';style.font.size=Pt(12);style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.space_after=Pt(6)
        style.paragraph_format.line_spacing=1.05
for level,size in [(1,16),(2,14),(3,12)]:
    st=next(x for x in doc.styles if x.style_id=='Heading'+str(level));st.font.size=Pt(size);st.font.bold=True
    st.paragraph_format.keep_with_next=True
doc.styles['Title'].font.size=Pt(22)
for p in doc.paragraphs:
    if 'GPU Sharing for Multi-Tenant Inference'==p.text.strip():p.style='Title'
for table in doc.tables:
    table.autofit=False
    pr=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        node=OxmlElement('w:'+edge);node.set(qn('w:val'),'single');node.set(qn('w:sz'),'4');node.set(qn('w:color'),'D9D9D9');borders.append(node)
    pr.append(borders)
    # Retain Pandoc's proportional column widths. Permit long rows to split naturally.
    for ri,row in enumerate(table.rows):
        if ri==0:
            hdr=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(hdr)
        for cell in row.cells:
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF2');cell._tc.get_or_add_tcPr().append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(4)
                p.paragraph_format.keep_with_next=False
                for run in p.runs:
                    run.font.size=Pt(12)
                    if ri==0:run.bold=True
for sec in doc.sections:
    p=sec.footer.paragraphs[0];p.alignment=1
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
out=report/'CEP_Task3_Integrated_First_Draft.docx';doc.save(out)
(report/'word_generation.json').write_text(json.dumps({'source':'report/main.tex and included TeX files','reference_order':'report/main.bbl','reference_count':len(refs),'tables':len(doc.tables),'paragraphs':len(doc.paragraphs),'same_claims_and_appendices':True,'output':out.name},indent=2))
print(out)
