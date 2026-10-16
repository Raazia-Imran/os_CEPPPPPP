"""Create the pre-Cycle2 human work handoff from committed Cycle1 counts."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parents[1]
p=json.loads((R/'evidence/review_progress.json').read_text())
d=Document();sec=d.sections[0];sec.top_margin=sec.bottom_margin=Inches(.7)
sec.left_margin=sec.right_margin=Inches(.75)
for name in ['Normal','Title','Heading 1','Heading 2']:
 s=d.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)

for style in d.styles:
 for border in list(style.element.xpath('.//w:pBdr')):
  border.getparent().remove(border)
d.styles['Normal'].font.size=Pt(11);d.styles['Normal'].paragraph_format.space_after=Pt(6)
d.styles['Title'].font.size=Pt(23);d.styles['Heading 1'].font.size=Pt(15)
d.add_paragraph('OS CEP remaining work and team actions','Title')
d.add_paragraph('Group B   Topic 23   Before Cycle 2   16 October 2026')
d.add_paragraph('The first recommendation pass is complete. The review is not yet finally validated. Cycle 2 will turn the existing evidence into an interim Task 2 and CEP report; it will not invent final inclusion decisions or independent reviewer votes.')
d.add_heading('What is already done',1)
t=d.add_table(rows=1,cols=2);t.autofit=False;t.columns[0].width=Inches(2.3);t.columns[1].width=Inches(4.4)
t.rows[0].cells[0].text='Work';t.rows[0].cells[1].text='Cycle 1 verified position'
rows=[('Imports and duplicates','2,291 occurrences; 44 applied duplicate links; 2,247 canonical records.'),('Preliminary screening','All 2,247 have recommendations: 1,181 exclude and 1,066 retain or uncertain. No first recommendations remain pending.'),('Final selection','Pending. Preliminary exclusions are not final exclusions.'),('Quality assessment','Five draft Q1-Q10 forms, containing 50 item scores on the Group B 0-2 scale.'),('Data extraction','26 outcome rows and five study forms. Unchecked fields still need extraction.'),('Citation and version checks','10 citation events; all 41 possible pairs audited, including five unresolved paper pairs.')]
for a,b in rows:
 c=t.add_row().cells;c[0].text=a;c[1].text=b
for i,row in enumerate(t.rows):
 for cell in row.cells:
  pr=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
  for edge in ['top','left','bottom','right']:
   e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
  pr.append(b)
  if i==0:
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'243B53');pr.append(sh)
   for run in cell.paragraphs[0].runs:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
  elif i%2:
   sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F2F5F8');pr.append(sh)
  for par in cell.paragraphs:par.paragraph_format.space_after=Pt(4)
d.add_heading('The part that needs a real person',1)
d.add_paragraph('An independent human review cannot be created by repeating the same first-review process. A teammate must read the evidence and record their own judgment. I can organize the queue, check sources and identify inconsistencies, but I cannot supply that teammate opinion or sign their name.')
d.add_paragraph('If missing search details exist only in a private account history or an unsupplied screenshot, the account holder must share the actual evidence. I can then extract and reconcile it. Do not reconstruct a date, query or hit count from memory or upload time.')
d.add_page_break()
d.add_heading('The exact human review assignments',1)
d.add_heading('Eshaal',2)
d.add_paragraph('Preserve Rameen\'s already completed ACM work. Supply or align its record IDs and original votes; do not repeat it as a new independent review. Start validation with the 148 priority candidate records in Human validation, then work through the remaining candidate queue. The complete queue has 1,066 records. Independently inspect the 127 sampled exclusions. Record vote, criterion, short reason, name and actual date; keep disagreements visible.')
d.add_paragraph('Recover the accepted ACM query, collection, field, filters, actual search date, displayed count and matching export pages if those records still exist. The claimed 1,600 hits and 1,603 exports do not reconcile to the two overlapping master files. The 79 numbered batches and 68 staged unmatched identifiers require reconciliation before counts change.')
d.add_heading('Rameen',2)
d.add_paragraph('Check reports GB-P001 to GB-P005: VastPipe, iGniter, gpulets, PipeSwitch and REEF conference. Use Report validation and Full-text criteria. Independently verify all 26 outcome IDs GB-O001 to GB-O026 against the exact page, table or figure. Check the hardware, workload, baseline, value, unit and statistic. Score all 50 Q1-Q10 items using 0, 1 or 2; justify N/A. Save your initial scores before comparing them with the draft.')
d.add_paragraph('Explicitly decide whether PipeSwitch\'s separable inference results satisfy the mixed training/inference rule, and verify REEF conference versus journal linkage. Never copy a Group A 0-1 score into this rubric.')
d.add_heading('Raazia',2)
d.add_paragraph('Coordinate the actual human judgments and resolve disagreements. Validate the three identity-link proposals, three companion-report proposals and five unresolved paper pairs. Accept only justified links and recalculate counts afterwards. Approve the eventual report-to-study mapping and final submission after evidence is reconciled.')
d.add_heading('Work I can continue doing',1)
d.add_paragraph('I can draft and check taxonomy, comparison tables, trade-offs, conflicts, RQ answers, conditional recommendations, gaps and validity; maintain the Excel and CSV registers; compile and inspect the PDF; verify bibliographic DOI metadata; locate available sources; and help complete remaining first-review extraction. Those tasks are not automatically manual just because they remain unfinished.')
d.add_heading('What remains after Cycle 2',1)
d.add_paragraph('Complete real validation and broader full-text eligibility, QA and extraction for every ultimately included study. Resolve search/version gaps and staged candidates. Complete citation chasing, freeze the eligible study set, derive final PRISMA and revise the interim conclusions. The instructor target is 15-20 eligible primary studies; five draft report checks do not establish that target. Finish with a reconciled DOI or official-source catalogue, final references, PDF layout check and submission.')
out=R/'task2/MANUAL_WORK_AND_NEXT_STEPS.docx';d.save(out)
print(out)
